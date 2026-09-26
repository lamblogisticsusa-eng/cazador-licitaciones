import io
import json
import logging
import os
import re
import sqlite3
import threading
import time
from datetime import datetime, timedelta, timezone

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from flask import Flask
from google import genai
from google.genai import types
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
import requests

from telegram import InlineKeyboardButton, InlineKeyboardMarkup, Update
from telegram.ext import (
    ApplicationBuilder,
    CallbackQueryHandler,
    CommandHandler,
    ContextTypes,
)

# ---------------------------------------------------------
# LOGGING
# ---------------------------------------------------------
logging.basicConfig(
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
    level=logging.INFO,
)
logger = logging.getLogger("Kiyomoto_Logistics")

# ---------------------------------------------------------
# SERVIDOR FLASK (Health Check para Render)
# ---------------------------------------------------------
flask_app = Flask(__name__)


@flask_app.route("/")
def health_check():
    return (
        "<h2>(✿◠‿◠) Kiyomoto Helper - L.A.M.B. Logistics LLC</h2>"
        "<p>Estado: Activa y rastreando SAM.gov 24/7 (Filtro COTS Estricto)</p>",
        200,
    )


def run_flask():
    port = int(os.environ.get("PORT", 10000))
    flask_app.run(host="0.0.0.0", port=port, use_reloader=False)


# ---------------------------------------------------------
# BASE DE DATOS SQLITE (Persistencia)
# ---------------------------------------------------------
DB_FILE = "kiyomoto_memory.db"


def init_db():
    conn = sqlite3.connect(DB_FILE)
    c = conn.cursor()
    c.execute(
        """
        CREATE TABLE IF NOT EXISTS notificadas (
            notice_id TEXT PRIMARY KEY,
            fecha TIMESTAMP
        )
    """
    )
    c.execute(
        """
        CREATE TABLE IF NOT EXISTS postuladas (
            solicitation_number TEXT PRIMARY KEY,
            fecha TIMESTAMP
        )
    """
    )
    conn.commit()
    conn.close()


def es_notificada(notice_id):
    conn = sqlite3.connect(DB_FILE)
    c = conn.cursor()
    c.execute("SELECT 1 FROM notificadas WHERE notice_id = ?", (notice_id,))
    res = c.fetchone()
    conn.close()
    return res is not None


def marcar_notificada(notice_id):
    conn = sqlite3.connect(DB_FILE)
    c = conn.cursor()
    c.execute(
        "INSERT OR IGNORE INTO notificadas VALUES (?, ?)",
        (notice_id, datetime.now(timezone.utc)),
    )
    conn.commit()
    conn.close()


def registrar_postulacion(sol_num):
    conn = sqlite3.connect(DB_FILE)
    c = conn.cursor()
    c.execute(
        "INSERT OR IGNORE INTO postuladas VALUES (?, ?)",
        (sol_num, datetime.now(timezone.utc)),
    )
    conn.commit()
    conn.close()


def obtener_postuladas():
    conn = sqlite3.connect(DB_FILE)
    c = conn.cursor()
    c.execute("SELECT solicitation_number FROM postuladas")
    rows = c.fetchall()
    conn.close()
    return [r[0] for r in rows]


def eliminar_postulacion(sol_num):
    conn = sqlite3.connect(DB_FILE)
    c = conn.cursor()
    c.execute("DELETE FROM postuladas WHERE solicitation_number = ?", (sol_num,))
    conn.commit()
    conn.close()


init_db()

# ---------------------------------------------------------
# VARIABLES DE ENTORNO Y IA
# ---------------------------------------------------------
TELEGRAM_BOT_TOKEN = os.getenv("TELEGRAM_BOT_TOKEN")
TELEGRAM_CHAT_ID = os.getenv("TELEGRAM_CHAT_ID")
SAM_API_KEY = os.getenv("SAM_API_KEY")
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")

MEMORIA_TEMPORAL = {}

ai_client = None
if GEMINI_API_KEY:
    try:
        ai_client = genai.Client(api_key=GEMINI_API_KEY)
        logger.info("Gemini API conectada con éxito (✿◠‿◠)")
    except Exception as e:
        logger.error(f"Error conectando Gemini API: {e}")


def escapar_markdown(texto: str) -> str:
    if not texto:
        return ""
    return (
        str(texto)
        .replace("_", "\\_")
        .replace("*", "\\*")
        .replace("`", "\\`")
        .replace("[", "\\[")         .replace("]", "\\]")
    )


# ---------------------------------------------------------
# GENERACIÓN DE DOCUMENTOS (DOCX & PDF)
# ---------------------------------------------------------
def generar_docx_rfq(lic_data):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    title_p = doc.add_paragraph()
    title_run = title_p.add_run("L.A.M.B. LOGISTICS LLC")
    title_run.bold = True
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run(
        "Government Contracting & Wholesale Supply Chain Division\n"
        "Email: logistics@lamblogistics.com"
    )
    sub_run.font.size = Pt(9)
    sub_run.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

    doc.add_paragraph("―" * 55)

    sol_num = lic_data.get("solicitation_number", "N/A")
    producto = lic_data.get("producto", "Commercial Product Specification")
    cantidad = lic_data.get("cantidad", 1)
    unidad = lic_data.get("unidad", "Units")
    zip_code = lic_data.get("zip_code", "USA")

    head_p = doc.add_paragraph()
    head_run = head_p.add_run("OFFICIAL REQUEST FOR QUOTATION (RFQ)")
    head_run.bold = True
    head_run.font.size = Pt(14)
    head_run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)

    doc.add_paragraph(f"Date: {datetime.now().strftime('%B %d, %Y')}")
    doc.add_paragraph(f"Solicitation Ref: {sol_num}\n")

    doc.add_paragraph(
        "Dear Commercial Quotations Department,\n\n"
        f"L.A.M.B. Logistics LLC is submitting a formal bid response for U.S. Federal Procurement Ref #{sol_num}. "
        "We kindly request your best wholesale unit pricing, availability, and lead time for the following items:\n"
    )

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Item Description"
    hdr[1].text = "Quantity"
    hdr[2].text = "Destination ZIP"

    row = table.add_row().cells
    row[0].text = str(producto)
    row[1].text = f"{cantidad:,} {unidad}"
    row[2].text = str(zip_code)

    doc.add_paragraph(
        "\nTerms & Requirements:\n"
        "1. Freight: Direct drop-shipment to specified destination ZIP.\n"
        "2. Payment Terms: Credit Card / Net 30.\n"
        "3. Outer Packaging: Must reference Solicitation Ref Number on packing slips.\n\n"
        "Please transmit your quote to logistics@lamblogistics.com at your earliest convenience."
    )

    p_sign = doc.add_paragraph(
        "\nSincerely,\n\nPurchasing & Procurement Team\nL.A.M.B. Logistics LLC"
    )
    p_sign.runs[0].bold = True

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    buf.name = f"RFQ_Vendor_{sol_num}.docx"
    return buf


def generar_pdf_po(lic_data):
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    w, h = letter

    c.setFont("Helvetica-Bold", 16)
    c.setFillColor(colors.HexColor("#1A365D"))
    c.drawString(50, h - 50, "L.A.M.B. LOGISTICS LLC")

    c.setFont("Helvetica", 9)
    c.setFillColor(colors.black)
    c.drawString(50, h - 65, "Government Contracting & Supply Chain Division")
    c.drawString(50, h - 77, "Email: logistics@lamblogistics.com")

    c.setFont("Helvetica-Bold", 14)
    c.setFillColor(colors.HexColor("#2B6CB0"))
    c.drawString(380, h - 50, "PURCHASE ORDER")

    c.setFont("Helvetica", 10)
    c.setFillColor(colors.black)
    c.drawString(380, h - 67, f"PO Date: {datetime.now().strftime('%Y-%m-%d')}")
    c.drawString(
        380,
        h - 80,
        f"Solicitation Ref: {lic_data.get('solicitation_number', 'N/A')}",
    )

    c.setStrokeColor(colors.HexColor("#CBD5E0"))
    c.setLineWidth(1)
    c.line(50, h - 95, w - 50, h - 95)

    c.setFont("Helvetica-Bold", 11)
    c.drawString(50, h - 120, "1. Item Specifications:")

    c.setFont("Helvetica", 10)
    y = h - 140
    c.drawString(65, y, f"• Product: {lic_data.get('producto', 'N/A')}")
    y -= 18
    c.drawString(
        65,
        y,
        f"• Quantity: {lic_data.get('cantidad', 1):,} {lic_data.get('unidad', 'Units')}",
    )
    y -= 18
    c.drawString(
        65, y, f"• Target Unit Cost Limit: ${lic_data.get('target_cost', 0):,.2f} USD"
    )
    y -= 18
    c.drawString(65, y, f"• Delivery Destination ZIP: {lic_data.get('zip_code', 'USA')}")

    y -= 35
    c.setFont("Helvetica-Bold", 11)
    c.drawString(50, y, "2. Fulfillment Terms:")
    c.setFont("Helvetica", 9)
    y -= 20
    for term in [
        "a. Direct drop-shipment required to final destination point.",
        "b. Packing slip must be attached without pricing references.",
        "c. Outer cartons must clearly display the Solicitation Number.",
        "d. Send tracking updates immediately to logistics@lamblogistics.com.",
    ]:
        c.drawString(65, y, term)
        y -= 15

    c.showPage()
    c.save()
    buf.seek(0)
    buf.name = f"Vendor_PO_{lic_data.get('solicitation_number', 'Doc')}.pdf"
    return buf


def generar_pdf_packing_list(lic_data):
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    w, h = letter

    c.setFont("Helvetica-Bold", 16)
    c.setFillColor(colors.HexColor("#1A365D"))
    c.drawString(50, h - 50, "PACKING LIST / SLIP")

    c.setFont("Helvetica", 10)
    c.setFillColor(colors.black)
    c.drawString(50, h - 67, "Prime Contractor: L.A.M.B. LOGISTICS LLC")
    c.drawString(50, h - 80, f"Date: {datetime.now().strftime('%Y-%m-%d')}")

    c.setStrokeColor(colors.HexColor("#CBD5E0"))
    c.line(50, h - 95, w - 50, h - 95)

    c.setFont("Helvetica-Bold", 11)
    c.drawString(50, h - 120, "Shipment Details:")

    c.setFont("Helvetica", 10)
    y = h - 140
    c.drawString(
        65, y, f"Solicitation Ref: {lic_data.get('solicitation_number', 'N/A')}"
    )
    y -= 18
    c.drawString(65, y, f"Item Description: {lic_data.get('producto', 'N/A')}")
    y -= 18
    c.drawString(
        65,
        y,
        f"Total Quantity: {lic_data.get('cantidad', 1):,} {lic_data.get('unidad', 'Units')}",
    )
    y -= 18
    c.drawString(65, y, f"Destination ZIP: {lic_data.get('zip_code', 'N/A')}")

    c.showPage()
    c.save()
    buf.seek(0)
    buf.name = f"Packing_List_{lic_data.get('solicitation_number', 'Doc')}.pdf"
    return buf


# ---------------------------------------------------------
# ANÁLISIS DE IA CON GEMINI (FILTRO COTS REFORZADO)
# ---------------------------------------------------------
def analizar_licitacion_ia(titulo, descripcion):
    if not ai_client:
        return {
            "es_producto_cots": True,
            "producto": titulo,
            "cantidad": 1,
            "unidad": "Unidades",
            "detalles": "Sin IA disponible.",
        }

    prompt = f"""
    Analiza esta licitación pública de compras públicas de EE. UU.:
    Título: {titulo}
    Descripción: {descripcion}

    REGLA CRÍTICA:
    Determina si la licitación es estrictamente para COMPRAR Y ENTREGAR PRODUCTOS FÍSICOS COMERCIALES (COTS) (ej. herramientas, suministros, piezas, equipos).
    Si se trata de SERVICIOS, TRABAJOS DE CAMPO, DEMOLICIÓN, REMOCIÓN, MANTENIMIENTO, CONSTRUCCIÓN O INSTALACIÓN EN SITIO, debes marcar "es_producto_cots": false.

    Responde ÚNICAMENTE en JSON con esta estructura exacta:
    {{
        "es_producto_cots": true/false,
        "producto": "Nombre claro del producto en español",
        "cantidad": 100,
        "unidad": "Unidades/Cajas/Kits",
        "detalles": "Resumen rápido de las especificaciones"
    }}
    """
    for attempt in range(2):
        try:
            resp = ai_client.models.generate_content(
                model="gemini-2.5-flash",
                contents=prompt,
                config=types.GenerateContentConfig(response_mime_type="application/json"),
            )
            text_clean = re.sub(
                r"```json\s*", "", resp.text, flags=re.IGNORECASE
            ).replace("```", "").strip()
            data = json.loads(text_clean)
            data["cantidad"] = max(1, int(data.get("cantidad", 1)))
            return data
        except Exception as e:
            logger.warning(f"Error parseando Gemini en intento {attempt + 1}: {e}")
            time.sleep(1)

    return {
        "es_producto_cots": True,
        "producto": titulo,
        "cantidad": 1,
        "unidad": "Unidades",
        "detalles": "Resumen no disponible.",
    }


def buscar_distribuidores_ia(producto, zip_code):
    if not ai_client:
        return [
            {
                "nombre": "Grainger Industrial",
                "tel": "(800) 472-4643",
                "web": "grainger.com",
            },
            {
                "nombre": "Fastenal Supply",
                "tel": "(800) 416-9400",
                "web": "fastenal.com",
            },
        ]

    prompt = f"""
    Encuentra 2 distribuidores o mayoristas reales en EE. UU. que vendan el producto '{producto}' y envíen al ZIP '{zip_code}'.
    Responde ÚNICAMENTE en JSON con una lista de 2 objetos conteniendo: "nombre", "tel", "web".
    """
    for attempt in range(2):
        try:
            resp = ai_client.models.generate_content(
                model="gemini-2.5-flash",
                contents=prompt,
                config=types.GenerateContentConfig(response_mime_type="application/json"),
            )
            text_clean = re.sub(
                r"```json\s*", "", resp.text, flags=re.IGNORECASE
            ).replace("```", "").strip()
            data = json.loads(text_clean)
            if isinstance(data, list) and len(data) > 0:
                return data
        except Exception as e:
            logger.warning(f"Error buscando proveedores en intento {attempt + 1}: {e}")
            time.sleep(1)

    return [
        {"nombre": "Grainger Supply", "tel": "(800) 472-4643", "web": "grainger.com"},
        {"nombre": "MSC Industrial Direct", "tel": "(800) 645-7270", "web": "mscdirect.com"},
    ]


# ---------------------------------------------------------
# CONEXIÓN SAM.GOV (FILTRO MEJORADO ANTI-SERVICIOS)
# ---------------------------------------------------------
def consultar_sam():
    if not SAM_API_KEY:
        return []

    hoy = datetime.now(timezone.utc)
    fecha_desde = (hoy - timedelta(days=10)).strftime("%m/%d/%Y")
    fecha_hasta = hoy.strftime("%m/%d/%Y")

    url = "https://api.sam.gov/prod/opportunities/v2/search"
    params = {
        "api_key": SAM_API_KEY,
        "postedFrom": fecha_desde,
        "postedTo": fecha_hasta,
        "limit": 100,
        "ptype": "k,o,p",
        "is_active": "true",
    }

    try:
        r = requests.get(url, params=params, timeout=25)
        if r.status_code == 200:
            opps = r.json().get("opportunitiesData", [])
        else:
            return []
    except Exception as e:
        logger.error(f"Error consultando SAM.gov: {e}")
        return []

    # Lista ampliada de exclusión de servicios y obras físicas
    excluir = [
        "construction", "service", "maintenance", "repair",
        "janitorial", "installation", "demolition", "removal",
        "disposal", "rental", "lease", "dredging", "painting",
        "inspection", "labor", "testing", "calibration", "renovation"
    ]
    candidatas = []

    for opp in opps:
        opp_id = opp.get("noticeId")
        if not opp_id or es_notificada(opp_id):
            continue

        titulo = opp.get("title", "")
        desc = opp.get("description", "")
        texto = f"{titulo} {desc}".lower()

        # Filtro 1: Palabras clave excluidas
        if any(kw in texto for kw in excluir):
            continue

        # Fechas y Cierre
        dias_restantes = 15
        cierre_str = "A la brevedad"
        deadline = opp.get("responseDeadLine")
        if deadline:
            try:
                dt_cierre = datetime.fromisoformat(deadline.replace("Z", "+00:00"))
                dias_restantes = (dt_cierre - hoy).days
                cierre_str = dt_cierre.strftime("%d/%m/%Y")
            except Exception:
                pass

        if dias_restantes < 1 or dias_restantes > 45:
            continue

        # Monto Estimado (Blindaje contra NoneType)
        monto_est = 35000.0
        award_dict = opp.get("award") or {}
        raw_amt = award_dict.get("amount")
        if raw_amt:
            try:
                monto_est = float(raw_amt)
            except (ValueError, TypeError):
                pass

        if monto_est > 250000.0:
            continue

        place_dict = opp.get("placeOfPerformance") or {}
        zip_code = place_dict.get("zip", "USA")

        candidatas.append(
            {
                "id": opp_id,
                "titulo": titulo or "Oportunidad COTS",
                "descripcion": desc,
                "solicitation_number": opp.get("solicitationNumber", "N/A"),
                "agencia": opp.get("department", "Agencia Federal"),
                "cierre_str": cierre_str,
                "dias_restantes": dias_restantes,
                "monto_est": monto_est,
                "zip_code": zip_code,
                "link": opp.get("uiLink", f"https://sam.gov/opp/{opp_id}/view"),
            }
        )

    candidatas.sort(key=lambda x: x["dias_restantes"])
    return candidatas[:5]


def verificar_adjudicaciones():
    postuladas = obtener_postuladas()
    if not SAM_API_KEY or not postuladas:
        return []

    hoy = datetime.now(timezone.utc)
    fecha_desde = (hoy - timedelta(days=14)).strftime("%m/%d/%Y")

    url = "https://api.sam.gov/prod/opportunities/v2/search"
    params = {
        "api_key": SAM_API_KEY,
        "postedFrom": fecha_desde,
        "postedTo": hoy.strftime("%m/%d/%Y"),
        "limit": 50,
        "ptype": "a",
    }

    try:
        r = requests.get(url, params=params, timeout=20)
        if r.status_code != 200:
            return []
        opps = r.json().get("opportunitiesData", [])
    except Exception:
        return []

    adjudicadas = []
    for opp in opps:
        sol_num = opp.get("solicitationNumber", "")
        if sol_num in postuladas:
            award = opp.get("award") or {}
            awardee = award.get("awardee") or {}
            adjudicadas.append(
                {
                    "solicitation_number": sol_num,
                    "titulo": opp.get("title", "Licitación"),
                    "monto": award.get("amount", "N/A"),
                    "ganador": awardee.get("name", "Desconocido"),
                    "link": opp.get(
                        "uiLink", f"https://sam.gov/opp/{opp.get('noticeId')}/view"
                    ),
                }
            )
    return adjudicadas


# ---------------------------------------------------------
# RASTREO Y NOTIFICACIÓN
# ---------------------------------------------------------
async def buscar_y_notificar(context: ContextTypes.DEFAULT_TYPE, target_chat_id=None):
    chat_id = target_chat_id or (context.job.chat_id if context.job else TELEGRAM_CHAT_ID)
    if not chat_id:
        return 0

    if isinstance(chat_id, str) and (chat_id.isdigit() or chat_id.startswith("-")):
        chat_id = int(chat_id)

    # 1. Notificar Adjudicaciones
    adjudicadas = verificar_adjudicaciones()
    for adj in adjudicadas:
        sol_num = adj["solicitation_number"]
        msg_adj = (
            "🏆 **¡RESULTADO DE ADJUDICACIÓN PUBLICADO!** 🏆\n\n"
            f"📋 **Solicitation #:** `{escapar_markdown(sol_num)}`\n"
            f"📌 **Título:** {escapar_markdown(adj['titulo'])}\n"
            f"💰 **Monto Adjudicado:** ${adj['monto']}\n"
            f"🏢 **Ganador:** {escapar_markdown(adj['ganador'])}\n\n"
            f"🔗 [Ver Registro en SAM.gov]({adj['link']})"
        )
        await context.bot.send_message(
            chat_id=chat_id, text=msg_adj, parse_mode="Markdown"
        )
        eliminar_postulacion(sol_num)

    # 2. Nuevas Licitaciones
    licitaciones = consultar_sam()
    enviadas_count = 0

    for lic in licitaciones:
        marcar_notificada(lic["id"])

        # Filtro 2: Validación por IA Gemini
        ia_res = analizar_licitacion_ia(lic["titulo"], lic["descripcion"])
        if not ia_res.get("es_producto_cots", True):
            logger.info(f"Omitiendo {lic['id']} por ser un servicio/obra física según Gemini.")
            continue

        cantidad = max(1, int(ia_res.get("cantidad", 1)))
        producto = ia_res.get("producto", lic["titulo"])
        unidad = ia_res.get("unidad", "Unidades")

        distribuidores = buscar_distribuidores_ia(producto, lic["zip_code"])

        monto_est = lic["monto_est"]
        precio_bid_unit = monto_est / cantidad
        costo_target_unit = (monto_est * 0.70) / cantidad

        flete = 400.0
        factoring = monto_est * 0.03
        costo_prov_total = monto_est * 0.70
        ganancia_est = monto_est - (costo_prov_total + flete + factoring)

        MEMORIA_TEMPORAL[lic["id"]] = {
            "solicitation_number": lic["solicitation_number"],
            "producto": producto,
            "cantidad": cantidad,
            "unidad": unidad,
            "zip_code": lic["zip_code"],
            "target_cost": costo_target_unit,
        }

        distrib_text = ""
        for idx, d in enumerate(distribuidores, 1):
            distrib_text += f"{idx}. *{escapar_markdown(d.get('nombre'))}* | Tel: {d.get('tel')} | Web: {d.get('web')}\n"

        mensaje = (
            "🚨 **NUEVA OPORTUNIDAD COTS DETECTADA** (✿◠‿◠)\n\n"
            f"📦 **Producto:** {escapar_markdown(producto)}\n"
            f"🔢 **Cantidad:** {cantidad:,} {unidad}\n"
            f"🏛 **Agencia:** {escapar_markdown(lic['agencia'])}\n"
            f"📍 **ZIP Entrega:** {lic['zip_code']}\n"
            f"📋 **Solicitation #:** `{escapar_markdown(lic['solicitation_number'])}`\n"
            f"⏳ **Cierre:** {lic['cierre_str']} ({lic['dias_restantes']} días restantes)\n\n"
            "💵 **ESTRATEGIA & MARGEN ESTIMADO**\n"
            f"• Presupuesto Est.: ${monto_est:,.2f} USD\n"
            f"• Precio Bid Sugerido: **${precio_bid_unit:,.2f} / unid.**\n"
            f"• Costo Máx. Compra Target: **<=${costo_target_unit:,.2f} / unid.**\n"
            f"📈 **Ganancia Neta Est.:** `${ganancia_est:,.2f} USD`\n\n"
            f"🏭 **PROVEEDORES POTENCIALES:**\n"
            f"{distrib_text}"
        )

        keyboard = InlineKeyboardMarkup(
            [
                [
                    InlineKeyboardButton("🔗 Abrir SAM.gov", url=lic["link"]),
                    InlineKeyboardButton(
                        "📄 Packing List", callback_data=f"pkg_{lic['id']}"
                    ),
                ],
                [
                    InlineKeyboardButton(
                        "📝 Generar PO (PDF)", callback_data=f"po_{lic['id']}"
                    ),
                    InlineKeyboardButton(
                        "✉️ Generar RFQ (DOCX)", callback_data=f"docx_{lic['id']}"
                    ),
                ],
            ]
        )

        await context.bot.send_message(
            chat_id=chat_id,
            text=mensaje,
            parse_mode="Markdown",
            reply_markup=keyboard,
        )
        enviadas_count += 1

    return enviadas_count


# ---------------------------------------------------------
# MANEJO DE COMANDOS Y BOTONES
# ---------------------------------------------------------
async def cmd_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    keyboard = InlineKeyboardMarkup(
        [
            [
                InlineKeyboardButton("🔎 Buscar Licitaciones", callback_data="btn_scan"),
                InlineKeyboardButton("🚀 Activar Monitoreo", callback_data="btn_on"),
            ],
            [
                InlineKeyboardButton("⏸ Pausar Monitoreo", callback_data="btn_off"),
                InlineKeyboardButton("📋 Mis Ofertas", callback_data="btn_postulaciones"),
            ],
        ]
    )

    saludo = (
        "¡Hola Bastian! (⺣◡⺣)♡*\n"
        "Soy **Kiyomoto**, lista para cazar las mejores oportunidades en SAM.gov para **L.A.M.B. Logistics LLC**.\n\n"
        "Presiona un botón abajo o usa los comandos rápidos:\n"
        "• `/scan` - Escanear oportunidades ahora mismo\n"
        "• `/on` - Activar rastreo automático cada hora\n"
        "• `/off` - Pausar rastreo automático\n"
        "• `/postulado <solicitation_num>` - Registrar oferta enviada\n"
        "• `/rfq <solicitation_num>` - Ver borrador rápido de correo\n"
        "• `/mis_postulaciones` - Ver licitaciones en seguimiento"
    )

    if update.message:
        await update.message.reply_text(
            saludo, reply_markup=keyboard, parse_mode="Markdown"
        )


async def cmd_scan(update: Update, context: ContextTypes.DEFAULT_TYPE):
    chat_id = update.effective_chat.id
    await context.bot.send_message(
        chat_id=chat_id,
        text="🔎 **Kiyomoto escaneando SAM.gov en tiempo real...** (✿◠‿◠)",
        parse_mode="Markdown",
    )
    encontradas = await buscar_y_notificar(context, target_chat_id=chat_id)
    if encontradas == 0:
        await context.bot.send_message(
            chat_id=chat_id,
            text="ℹ️ **Sin novedades:** No se encontraron licitaciones nuevas de productos COTS en este momento.",
            parse_mode="Markdown",
        )


async def cmd_on(update: Update, context: ContextTypes.DEFAULT_TYPE):
    chat_id = update.effective_chat.id if update.effective_chat else update.callback_query.message.chat_id
    job_name = f"job_monitor_{chat_id}"

    if not context.job_queue:
        msg = "❌ Error: APScheduler no configurado correctamente."
        await context.bot.send_message(chat_id=chat_id, text=msg)
        return

    if context.job_queue.get_jobs_by_name(job_name):
        await context.bot.send_message(
            chat_id=chat_id,
            text="🟢 El monitoreo automático **ya está activo**.",
            parse_mode="Markdown",
        )
        return

    context.job_queue.run_repeating(
        buscar_y_notificar, interval=3600, first=1, chat_id=chat_id, name=job_name
    )
    await context.bot.send_message(
        chat_id=chat_id,
        text="🚀 **Monitoreo Automático ACTIVADO:** Rastreo cada 60 minutos (⺣◡⺣)♡*",
        parse_mode="Markdown",
    )


async def cmd_off(update: Update, context: ContextTypes.DEFAULT_TYPE):
    chat_id = update.effective_chat.id if update.effective_chat else update.callback_query.message.chat_id
    job_name = f"job_monitor_{chat_id}"

    if context.job_queue:
        jobs = context.job_queue.get_jobs_by_name(job_name)
        for j in jobs:
            j.schedule_removal()

    await context.bot.send_message(
        chat_id=chat_id,
        text="⏸ **Monitoreo PAUSADO:** Kiyomoto descansará por ahora.",
        parse_mode="Markdown",
    )


async def cmd_postulado(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not context.args:
        await update.message.reply_text(
            "⚠️ **Uso:** `/postulado SPE8E6-26-U-0012`", parse_mode="Markdown"
        )
        return

    sol_num = context.args[0].strip()
    registrar_postulacion(sol_num)
    await update.message.reply_text(
        f"✅ **Licitación registrada:** `{escapar_markdown(sol_num)}`\nTe avisaré en cuanto SAM.gov publique la adjudicación.",
        parse_mode="Markdown",
    )


async def cmd_mis_postulaciones(update: Update, context: ContextTypes.DEFAULT_TYPE):
    postuladas = obtener_postuladas()
    if not postuladas:
        await update.message.reply_text(
            "ℹ️ No tienes licitaciones registradas bajo seguimiento."
        )
        return

    lista = "\n".join([f"• `{escapar_markdown(num)}`" for num in postuladas])
    await update.message.reply_text(
        f"📋 **Licitaciones Bajo Monitoreo:**\n\n{lista}", parse_mode="Markdown"
    )


async def cmd_rfq(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not context.args:
        await update.message.reply_text("⚠️ **Uso:** `/rfq <solicitation_number>`")
        return

    sol_num = context.args[0].strip()
    mensaje_rfq = (
        "📧 **TEMPLATE DE CORREO RFQ**\n\n"
        "```text\n"
        f"Subject: Urgent RFQ Quote Request - L.A.M.B. Logistics LLC (Ref: {sol_num})\n\n"
        "Dear Sales Department,\n\n"
        f"L.A.M.B. Logistics LLC is preparing a formal offer for U.S. Federal Solicitation {sol_num}.\n\n"
        "Could you please share your wholesale unit price, availability, and delivery lead time?\n\n"
        "Key Delivery Terms:\n"
        "• Direct Freight: Destination ZIP specified in purchase order\n"
        "• Payment: Credit Card / Net 30\n"
        "• Reference: Include Solicitation Number on shipping paperwork\n\n"
        "Please transmit your quote to logistics@lamblogistics.com.\n\n"
        "Best Regards,\n"
        "Purchasing Team\n"
        "L.A.M.B. Logistics LLC\n"
        "```"
    )
    await update.message.reply_text(mensaje_rfq, parse_mode="Markdown")


async def boton_callback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    data = query.data

    if data == "btn_scan":
        await cmd_scan(update, context)

    elif data == "btn_on":
        await cmd_on(update, context)

    elif data == "btn_off":
        await cmd_off(update, context)

    elif data == "btn_postulaciones":
        await cmd_mis_postulaciones(update, context)

    elif data.startswith("pkg_"):
        lic_id = data.replace("pkg_", "")
        lic_data = MEMORIA_TEMPORAL.get(
            lic_id,
            {"solicitation_number": lic_id, "producto": "Producto COTS", "cantidad": 1},
        )
        pdf = generar_pdf_packing_list(lic_data)
        await context.bot.send_document(
            chat_id=query.message.chat_id,
            document=pdf,
            caption="📄 **Packing List (PDF)**",
        )

    elif data.startswith("po_"):
        lic_id = data.replace("po_", "")
        lic_data = MEMORIA_TEMPORAL.get(
            lic_id,
            {"solicitation_number": lic_id, "producto": "Producto COTS", "cantidad": 1},
        )
        pdf = generar_pdf_po(lic_data)
        await context.bot.send_document(
            chat_id=query.message.chat_id,
            document=pdf,
            caption="📝 **Vendor Purchase Order (PDF)**",
        )

    elif data.startswith("docx_"):
        lic_id = data.replace("docx_", "")
        lic_data = MEMORIA_TEMPORAL.get(
            lic_id,
            {"solicitation_number": lic_id, "producto": "Producto COTS", "cantidad": 1},
        )
        docx = generar_docx_rfq(lic_data)
        await context.bot.send_document(
            chat_id=query.message.chat_id,
            document=docx,
            caption="✉️ **Official RFQ Document (DOCX)**",
        )


# ---------------------------------------------------------
# PUNTO DE ENTRADA PRINCIPAL
# ---------------------------------------------------------
def main():
    flask_thread = threading.Thread(target=run_flask, daemon=True)
    flask_thread.start()

    if not TELEGRAM_BOT_TOKEN:
        logger.error("TELEGRAM_BOT_TOKEN no configurado.")
        return

    app = ApplicationBuilder().token(TELEGRAM_BOT_TOKEN).build()

    # Comandos
    app.add_handler(CommandHandler("start", cmd_start))
    app.add_handler(CommandHandler("scan", cmd_scan))
    app.add_handler(CommandHandler("on", cmd_on))
    app.add_handler(CommandHandler("off", cmd_off))
    app.add_handler(CommandHandler("postulado", cmd_postulado))
    app.add_handler(CommandHandler("mis_postulaciones", cmd_mis_postulaciones))
    app.add_handler(CommandHandler("rfq", cmd_rfq))
    app.add_handler(CallbackQueryHandler(boton_callback))

    # Tarea automática si existe CHAT_ID
    if TELEGRAM_CHAT_ID and app.job_queue:
        try:
            chat_target = (
                int(TELEGRAM_CHAT_ID)
                if str(TELEGRAM_CHAT_ID).lstrip("-").isdigit()
                else TELEGRAM_CHAT_ID
            )
            app.job_queue.run_repeating(
                buscar_y_notificar,
                interval=3600,
                first=10,
                chat_id=chat_target,
                name=f"job_monitor_{chat_target}",
            )
            logger.info("Programador de tareas listo (✿◠‿◠)")
        except Exception as e:
            logger.error(f"Error en programador de tareas: {e}")

    logger.info("Kiyomoto Helper iniciada y lista.")
    
    app.run_polling(drop_pending_updates=True)


if __name__ == "__main__":
    main()
